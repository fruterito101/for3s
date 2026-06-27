"""Tests de los módulos nuevos del pulido del MVP (sesiones 2026-06-15 a 18).

LÓGICA PURA — NO llaman a GitHub/Claude/red. Cubren: categorización y reparto
de archivos (subbloques), conversor MD→HTML (md_html), zonas horarias (tiempo),
parseo HTML (web_fetch), y los detectores org/repo/modo (conversation).
"""


from for3s_core import multimodal, subbloques, tiempo, tool_loop, web_fetch
from for3s_core.conversation import extraer_org, extraer_owner_repo, huele_a_github
from for3s_core.md_html import md_a_html_telegram


# ───────────────────────── subbloques: categorización ─────────────────────────
class TestCategoria:
    def test_readme_raiz_es_categoria_propia(self):
        assert subbloques._categoria("README.md") == "readme"
        assert subbloques._categoria("readme") == "readme"

    def test_readme_de_subcarpeta_es_doc(self):
        assert subbloques._categoria("docs/README.md") == "doc"

    def test_codigo_fuente(self):
        for f in ("src/app/page.tsx", "lib/data.ts", "main.py", "server.go"):
            assert subbloques._categoria(f) == "src", f

    def test_config(self):
        for f in ("package.json", "tsconfig.json", "Dockerfile", ".github/workflows/ci.yml"):
            assert subbloques._categoria(f) == "config", f

    def test_doc(self):
        assert subbloques._categoria("CHANGELOG.md") == "doc"
        assert subbloques._categoria("docs/guia.md") == "doc"

    def test_test(self):
        assert subbloques._categoria("tests/test_foo.py") == "test"
        assert subbloques._categoria("src/utils.test.ts") == "test"

    def test_otro(self):
        assert subbloques._categoria("LICENSE") == "otro"


# ───────────────────────── subbloques: capas de ejecución ─────────────────────
class TestCapaEjecucion:
    def test_entry_point_es_capa_0(self):
        for f in ("app/page.tsx", "app/layout.tsx", "main.py", "index.js", "server.ts"):
            assert subbloques._capa_ejecucion(f) == 0, f

    def test_data_es_capa_1(self):
        assert subbloques._capa_ejecucion("lib/data.ts") == 1
        assert subbloques._capa_ejecucion("src/store/user.ts") == 1

    def test_componentes_alto_capa_2(self):
        assert subbloques._capa_ejecucion("components/sections/Hero.tsx") == 2

    def test_primitivos_capa_3(self):
        assert subbloques._capa_ejecucion("components/ui/button.tsx") == 3
        assert subbloques._capa_ejecucion("lib/utils.ts") == 3

    def test_orden_afuera_adentro(self):
        # entry < data < componentes < primitivos
        c = subbloques._capa_ejecucion
        assert c("app/page.tsx") < c("lib/data.ts")
        assert c("lib/data.ts") < c("components/sections/X.tsx") < c("lib/utils.ts")


# ───────────────────────── subbloques: reparto por categoría ──────────────────
class TestReparto:
    archivos = (
        ["README.md"]
        + [f"docs/d{i}.md" for i in range(30)]          # 30 docs
        + ["package.json", "tsconfig.json"]              # 2 config
        + [f"src/c{i}.ts" for i in range(50)]            # 50 src
        + [f"tests/t{i}.test.ts" for i in range(20)]     # 20 test
        + [f"x{i}.bin" for i in range(10)]               # 10 otro
    )

    def test_simple_lee_pocos(self):
        sel, cob = subbloques.repartir_por_categoria(self.archivos, profundo=False)
        # SIMPLE: doc<=5, src<=8 (cuotas de CUOTA_CATEGORIA_SIMPLE)
        assert cob["doc"]["leidos"] <= 5
        assert cob["src"]["leidos"] <= 8
        assert cob["readme"]["leidos"] == 1  # README completo

    def test_profundo_lee_mucho_mas(self):
        sel, cob = subbloques.repartir_por_categoria(self.archivos, profundo=True)
        # PROFUNDO: src casi todo (cuota 120 > 50), doc hasta 25
        assert cob["src"]["leidos"] == 50
        assert cob["doc"]["leidos"] == 25  # tope 25 de los 30

    def test_profundo_lee_mas_src_que_simple(self):
        _, cs = subbloques.repartir_por_categoria(self.archivos, profundo=False)
        _, cp = subbloques.repartir_por_categoria(self.archivos, profundo=True)
        assert cp["src"]["leidos"] > cs["src"]["leidos"]

    def test_readme_primero_en_la_seleccion(self):
        sel, _ = subbloques.repartir_por_categoria(self.archivos, profundo=True)
        assert sel[0] == "README.md"  # README va PRIMERO

    def test_recencia_prioriza(self):
        archivos = ["src/a.ts", "src/b.ts", "src/c.ts"]
        recencia = {"src/c.ts": 0, "src/a.ts": 1}  # c más reciente que a
        sel, _ = subbloques.repartir_por_categoria(archivos, recencia, profundo=True)
        # c (más reciente) antes que a; b (sin recencia) al final
        assert sel.index("src/c.ts") < sel.index("src/a.ts") < sel.index("src/b.ts")


# ───────────────────────── md_html: conversor ─────────────────────────────────
class TestMdHtml:
    def test_negrita(self):
        assert "<b>x</b>" in md_a_html_telegram("**x**")

    def test_codigo_inline(self):
        assert "<code>foo()</code>" in md_a_html_telegram("`foo()`")

    def test_bloque_codigo_a_pre(self):
        out = md_a_html_telegram("```python\nx < 5\n```")
        assert "<pre>" in out and "&lt;" in out  # < escapado dentro de pre

    def test_escapa_html_peligroso(self):
        out = md_a_html_telegram("<script>evil</script>")
        assert "&lt;script&gt;" in out and "<script>" not in out

    def test_header_a_negrita(self):
        assert "<b>Título</b>" in md_a_html_telegram("# Título")

    def test_enlace(self):
        out = md_a_html_telegram("[txt](https://x.com)")
        assert '<a href="https://x.com">txt</a>' in out

    def test_tags_balanceados(self):
        import re
        out = md_a_html_telegram("**a** y `b` y *c* y # H\n```\nx\n```")
        for tag in ("b", "code", "pre", "i"):
            assert len(re.findall(f"<{tag}>", out)) == len(re.findall(f"</{tag}>", out)), tag


# ───────────────────────── tiempo: zonas ──────────────────────────────────────
class TestTiempo:
    def test_zona_con_pais(self):
        assert tiempo.zona_de_language("es-MX") == "America/Mexico_City"
        assert tiempo.zona_de_language("es-AR") == "America/Argentina/Buenos_Aires"
        assert tiempo.zona_de_language("en-US") == "America/New_York"

    def test_zona_sin_pais_default_cdmx(self):
        assert tiempo.zona_de_language("es") == tiempo.ZONA_DEFAULT
        assert tiempo.zona_de_language("en") == tiempo.ZONA_DEFAULT

    def test_zona_none_default(self):
        assert tiempo.zona_de_language(None) == tiempo.ZONA_DEFAULT

    def test_contexto_temporal_menciona_usuario_no_servidor(self):
        txt = tiempo.contexto_temporal("es-MX")
        assert "USUARIO" in txt and "no del servidor" in txt
        assert "America/Mexico_City" in txt


# ───────────────────────── web_fetch: parseo HTML ─────────────────────────────
class TestWebFetch:
    def test_extrae_titulo(self):
        html = "<html><head><title>Mi Página</title></head><body>hola</body></html>"
        out = web_fetch._html_a_texto(html)
        assert "Mi Página" in out

    def test_quita_scripts(self):
        out = web_fetch._html_a_texto("<body><script>malo()</script>texto real</body>")
        assert "malo()" not in out and "texto real" in out

    def test_extrae_og(self):
        html = '<meta property="og:title" content="Evento X">'
        out = web_fetch._html_a_texto(html)
        assert "Evento X" in out

    def test_largo_contenido_separa_cabecera(self):
        # solo cuenta lo que va tras CONTENIDO:, no el TÍTULO/DESC
        txt = "TÍTULO: x\nDESCRIPCIÓN: y\n\nCONTENIDO:\nhola mundo"
        assert web_fetch._largo_contenido(txt) == len("hola mundo")

    def test_largo_contenido_sin_marcador_mide_todo(self):
        assert web_fetch._largo_contenido("abcde") == 5

    def test_huele_a_login_detecta(self):
        assert web_fetch._huele_a_login("Please Sign in to continue")
        assert web_fetch._huele_a_login("Debes iniciar sesión")

    def test_huele_a_login_pagina_normal_no(self):
        assert not web_fetch._huele_a_login("Bienvenido al blog de ejemplo")

    def test_normaliza_ignora_esquema_y_slash(self):
        assert web_fetch._normaliza("https://X.com/") == web_fetch._normaliza("http://x.com")

    def test_normaliza_detecta_destino_distinto(self):
        assert web_fetch._normaliza("https://a.co/d/x") != web_fetch._normaliza("https://amazon.com/dp/y")

    def test_huele_a_antibot_amazon(self):
        assert web_fetch._huele_a_antibot("Click the button below to continue shopping")

    def test_huele_a_antibot_cloudflare(self):
        assert web_fetch._huele_a_antibot("Checking your browser before accessing")

    def test_antibot_pagina_normal_no(self):
        assert not web_fetch._huele_a_antibot("Bienvenido, este es un blog normal")


# ───────────────────────── conversation: detectores ───────────────────────────
class TestDetectores:
    def test_repo_completo(self):
        r = extraer_owner_repo("https://github.com/octocat/Hello-World")
        assert r == ("octocat", "Hello-World")

    def test_pr_no_es_repo_completo(self):
        assert extraer_owner_repo("https://github.com/foo/bar/pull/3") is None

    def test_org_sin_repo(self):
        assert extraer_org("https://github.com/All-Hands-AI analiza") == "All-Hands-AI"

    def test_org_none_si_es_repo(self):
        assert extraer_org("https://github.com/foo/bar") is None

    def test_huele_a_github(self):
        assert huele_a_github("analiza https://github.com/x/y")
        assert huele_a_github("cuántos ISSUES tiene godinez-studio")

    def test_no_huele_a_github(self):
        assert not huele_a_github("hola cómo estás")
        assert not huele_a_github("qué hora es")

    def test_url_web_no_es_github(self):
        # fix falso positivo (2026-06-19): dominio.com/path NO es repo
        assert not huele_a_github(
            "https://www.tvazteca.com/aztecadeportes/mundial-2026/envivo el partido")
        assert not huele_a_github("resume https://react.dev/learn/thinking-in-react")
        assert not huele_a_github("https://luma.com/yov2v22b que es esto")
        assert not huele_a_github("checa https://ethglobal.com/events/2026")

    def test_github_sigue_detectandose_con_url(self):
        # NO romper: github real + owner/repo en texto humano siguen detectándose
        assert huele_a_github("analiza https://github.com/cli/cli")
        assert huele_a_github("comenta en github.com/fruterito101/Proyecto")
        assert huele_a_github("cuantos PRs tiene acme/demo-repo")


# ───────────────────────── multimodal: adjuntos ───────────────────────────────
class TestMultimodal:
    def test_imagen_por_mime(self):
        bloques = multimodal.procesar_adjunto(b"\x89PNG fake", nombre="x", mime="image/png")
        assert bloques[0]["type"] == "image"
        assert bloques[0]["source"]["media_type"] == "image/png"

    def test_imagen_por_extension(self):
        bloques = multimodal.procesar_adjunto(b"fake", nombre="foto.JPG", mime="")
        assert bloques[0]["type"] == "image"
        assert bloques[0]["source"]["media_type"] == "image/jpeg"

    def test_pdf_es_document(self):
        bloques = multimodal.procesar_adjunto(b"%PDF-1.4 fake", nombre="x.pdf",
                                              mime="application/pdf")
        assert bloques[0]["type"] == "document"
        assert bloques[0]["source"]["media_type"] == "application/pdf"

    def test_word_extrae_texto(self):
        import io

        from docx import Document
        doc = Document()
        doc.add_paragraph("Hola desde Word")
        buf = io.BytesIO()
        doc.save(buf)
        bloques = multimodal.procesar_adjunto(buf.getvalue(), nombre="d.docx", mime="")
        assert bloques[0]["type"] == "text"
        assert "Hola desde Word" in bloques[0]["text"]

    def test_excel_extrae_celdas(self):
        import io

        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["nombre", "edad"])
        ws.append(["Ana", 30])
        buf = io.BytesIO()
        wb.save(buf)
        bloques = multimodal.procesar_adjunto(buf.getvalue(), nombre="h.xlsx", mime="")
        assert bloques[0]["type"] == "text"
        assert "Ana" in bloques[0]["text"] and "edad" in bloques[0]["text"]

    def test_tipo_no_soportado(self):
        import pytest
        with pytest.raises(multimodal.ArchivoNoSoportado):
            multimodal.procesar_adjunto(b"zip", nombre="a.zip", mime="application/zip")

    def test_archivo_vacio(self):
        import pytest
        with pytest.raises(multimodal.ArchivoNoSoportado):
            multimodal.procesar_adjunto(b"", nombre="x.png", mime="image/png")

    def test_archivo_demasiado_grande(self):
        import pytest
        enorme = b"x" * (multimodal.MAX_BYTES + 1)
        with pytest.raises(multimodal.ArchivoNoSoportado):
            multimodal.procesar_adjunto(enorme, nombre="x.png", mime="image/png")

    def test_descripcion_corta(self):
        assert multimodal.descripcion_corta("x.pdf", "") == "PDF"
        assert multimodal.descripcion_corta("f.png", "") == "imagen"

    def test_pdf_gigante_avisa_honesto(self):
        # PDF > MAX_BYTES_NATIVO (8MB) → aviso honesto, no intentar y morir
        import pytest
        grande = b"%PDF-1.4" + b"x" * (multimodal.MAX_BYTES_NATIVO + 1)
        with pytest.raises(multimodal.ArchivoNoSoportado):
            multimodal.procesar_adjunto(grande, nombre="big.pdf", mime="application/pdf")

    def test_pdf_normal_si_pasa(self):
        # un PDF chico (bajo el umbral nativo) sí se procesa
        bloques = multimodal.procesar_adjunto(b"%PDF-1.4 chico", nombre="ok.pdf",
                                              mime="application/pdf")
        assert bloques[0]["type"] == "document"


# ───────────────────────── tool_loop: conteo exacto ───────────────────────────
class TestConteo:
    def test_search_tools_en_whitelist(self):
        # las tools de CONTEO exacto (total_count en 1 llamada) deben estar
        # disponibles para el agente, si no los conteos grandes quedan parciales.
        assert "search_pull_requests" in tool_loop.MVP_TOOLS
        assert "search_issues" in tool_loop.MVP_TOOLS

    def test_list_tools_siguen_disponibles(self):
        # no rompimos las de listar (siguen sirviendo para 'los últimos N')
        assert "list_pull_requests" in tool_loop.MVP_TOOLS
        assert "list_issues" in tool_loop.MVP_TOOLS

    def test_max_tool_rounds_no_subio(self):
        # decisión de diseño: NO subir las vueltas (dispara rate-limit); el fix es la
        # herramienta correcta (search_*), no más iteraciones.
        assert tool_loop.MAX_TOOL_ROUNDS == 5


# ───────────────────────── write tools: seguridad ─────────────────────────────
class TestWriteTools:
    def test_solo_4_write_permitidas(self):
        from for3s_core import tool_loop
        assert tool_loop.WRITE_TOOLS_PERMITIDAS == {
            "add_issue_comment", "create_issue",
            "create_pull_request", "create_pull_request_review",
        }

    def test_destructive_NO_permitidas(self):
        # la garantía central: NUNCA estas en la whitelist (ni schema inyectado)
        from for3s_core import tool_loop
        prohibidas = {"merge_pull_request", "delete_repository",
                      "create_repository", "push_files", "create_or_update_file",
                      "update_pull_request_branch"}
        assert not (prohibidas & tool_loop.WRITE_TOOLS_PERMITIDAS)
        nombres_schema = {s["name"] for s in tool_loop.WRITE_TOOL_SCHEMAS}
        assert not (prohibidas & nombres_schema)

    def test_schemas_cubren_las_permitidas(self):
        from for3s_core import tool_loop
        nombres_schema = {s["name"] for s in tool_loop.WRITE_TOOL_SCHEMAS}
        assert nombres_schema == tool_loop.WRITE_TOOLS_PERMITIDAS

    def test_write_no_estan_en_read_whitelist(self):
        # las write NUNCA deben colarse en MVP_TOOLS (read) → no se auto-ejecutan
        from for3s_core import tool_loop
        assert not (tool_loop.WRITE_TOOLS_PERMITIDAS & tool_loop.MVP_TOOLS)

    def test_directive_menciona_confirmacion(self):
        from for3s_core.conversation import TOOL_DIRECTIVE
        assert "confirmaci" in TOOL_DIRECTIVE.lower()
        assert "create_pull_request" in TOOL_DIRECTIVE

    def test_intencion_escritura_vs_analisis(self):
        # routing fix (2026-06-18): 'comenta/crea' con URL de repo → write,
        # 'analiza' con URL de repo → análisis. Replica la lógica del canal.
        from for3s_core.text_normalize import normalizar
        PAL = ("comenta", "comentar", "crea un issue", "crear issue",
               "crea issue", "abre un issue", "crea un pr", "review")
        def quiere_escribir(m):
            t = normalizar(m)
            return any(p in t for p in PAL)
        assert quiere_escribir("comenta hola en el issue #1 de github.com/o/r")
        assert quiere_escribir("crea un issue en github.com/o/r")
        assert not quiere_escribir("analiza https://github.com/cli/cli")
        assert not quiere_escribir("cuantos PRs tiene cli/cli")


# ───────────────────────── cache Valkey: lógica pura ──────────────────────────
class TestCache:
    def test_ttl_por_tool(self):
        from for3s_core.cache import GitHubCache
        assert GitHubCache.cacheable("get_file_contents") == 300
        assert GitHubCache.cacheable("list_issues") == 30
        assert GitHubCache.cacheable("search_code") == 900

    def test_write_tools_NO_cacheables(self):
        # garantía: las write NUNCA se cachean (devuelven None = no cache)
        from for3s_core.cache import GitHubCache
        for w in ("add_issue_comment", "create_issue",
                  "create_pull_request", "create_pull_request_review"):
            assert GitHubCache.cacheable(w) is None, w

    def test_never_cache(self):
        from for3s_core.cache import GitHubCache
        assert GitHubCache.cacheable("get_pull_request_status") is None
        assert GitHubCache.cacheable("get_pull_request_files") is None

    def test_tool_desconocida_no_cachea(self):
        from for3s_core.cache import GitHubCache
        assert GitHubCache.cacheable("tool_inventada") is None

    def test_key_estable_sin_importar_orden_args(self):
        from for3s_core.cache import GitHubCache
        k1 = GitHubCache._key("brian", "issue_read", {"owner": "x", "repo": "y", "n": 1})
        k2 = GitHubCache._key("brian", "issue_read", {"n": 1, "repo": "y", "owner": "x"})
        assert k1 == k2

    def test_key_distinta_por_workspace(self):
        # multi-tenant: dos workspaces NUNCA comparten cache
        from for3s_core.cache import GitHubCache
        ka = GitHubCache._key("brian", "issue_read", {"n": 1})
        kb = GitHubCache._key("otro", "issue_read", {"n": 1})
        assert ka != kb

    def test_key_distinta_por_args(self):
        from for3s_core.cache import GitHubCache
        ka = GitHubCache._key("brian", "issue_read", {"n": 1})
        kb = GitHubCache._key("brian", "issue_read", {"n": 2})
        assert ka != kb
